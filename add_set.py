#!/usr/bin/env python3
"""नया सेट वेबसाइट और Word फ़ाइल में जोड़ें:
    python3 tools/add_set.py NN tools/src/setNN.txt "15 अक्टूबर 2026"

यह स्क्रिप्ट:
  1. setNN.txt पार्स और सत्यापित करती है (100 प्रश्न, खंड 15/15/20/10/15/5/5/5/5/5, चार विकल्प,
     एक उत्तर, व्याख्या, उत्तर-वितरण 22–28 प्रति अक्षर, एक ही अक्षर लगातार 4 बार नहीं)
  2. cg_lecturer_setNN.json लिखती है
  3. index.html के DATA.cg.sets में सेट जोड़ती/बदलती है और शीर्षक की सेट-सीमा अद्यतन करती है
  4. README.md की "नवीनतम" पंक्ति अद्यतन करती है
  5. CG_Lecturer_History_Set_NN_Syllabus_Based_DD_Mon_YYYY.docx बनाती है (उत्तरमाला नए पृष्ठ से)
  6. tools/lists/setNN.md और tools/log.json अद्यतन करती है
दोबारा चलाने पर वही सेट बदल दिया जाता है (दोहरी प्रविष्टि नहीं बनती)।
"""
import json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (ROOT, SECTIONS, LETTERS, DIFF, DIFF_NAME, parse_src, validate,
                    read_index, write_index, hindi_date)

PART = {0: 'भाग अ', 1: 'भाग अ', 2: 'भाग अ', 3: 'भाग अ', 4: 'भाग ब'}


def build_docx(no, qs, date_txt, path, n_old):
    from docx import Document
    from docx.shared import Pt, Emu, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Emu(7560310), Emu(10692130)
    sec.left_margin = sec.right_margin = Emu(720090)
    sec.top_margin = sec.bottom_margin = Emu(685800)
    st = doc.styles['Normal']
    st.font.name = 'Nirmala UI'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:cs'), 'Nirmala UI')
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Nirmala UI')
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), '21'); st.element.rPr.append(szcs)
    st.paragraph_format.space_after = Pt(2)

    def para(text='', bold=None, size=None, align=None, after=None, before=None, keep=False):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text); r.bold = bold
            if size: r.font.size = Pt(size)
        if align: p.alignment = align
        if after is not None: p.paragraph_format.space_after = Pt(after)
        if before is not None: p.paragraph_format.space_before = Pt(before)
        if keep: p.paragraph_format.keep_with_next = True
        return p

    para(f'छत्तीसगढ़ व्याख्याता (इतिहास) भर्ती परीक्षा — मॉक टेस्ट सेट {no}', True, 17, WD_ALIGN_PARAGRAPH.CENTER, 4)
    para(f'SCERT छत्तीसगढ़ पाठ्यक्रम पर आधारित  |  कुल 100 प्रश्न / 100 अंक  |  {date_txt}', True, 12, WD_ALIGN_PARAGRAPH.CENTER, 2)
    para('भाग अ: प्राचीन 15, मध्यकालीन 15, आधुनिक (स्वातंत्र्योत्तर व छत्तीसगढ़ सहित) 20, यूरोपीय 10  |  '
         'भाग ब: शैक्षिक मनोविज्ञान, आकलन-मूल्यांकन, शिक्षणशास्त्र, शैक्षिक अभिवृत्ति 15  |  '
         'भाग स: हिंदी 5, English 5, मानसिक योग्यता 5, कंप्यूटर 5, सामान्य ज्ञान 5', None, 10, WD_ALIGN_PARAGRAPH.CENTER, 8)
    para('निर्देश:', True, 10, after=1)
    for t in ['1. सभी 100 प्रश्न वस्तुनिष्ठ हैं; प्रत्येक के चार विकल्पों में से केवल एक सही है।',
              '2. खंडवार प्रश्न-संख्या SCERT पाठ्यक्रम के अंक-विभाजन के अनुसार है (प्रति प्रश्न 1 अंक माना गया है)। '
              'समय और नकारात्मक अंकन आधिकारिक परीक्षा-सूचना के अनुसार मानें।',
              '3. उत्तरमाला और संक्षिप्त व्याख्या प्रश्नपत्र के बाद नए पृष्ठ से दी गई है।',
              f'4. प्रश्न वेबसाइट और पिछले सभी सेटों के {n_old:,} प्रश्नों से दोहराव-जाँच के बाद चुने गए हैं। '
              'यह अभ्यास-पत्र है, पिछले वर्ष का आधिकारिक प्रश्नपत्र नहीं।']:
        r = para(t, False, 10, after=1)
    groups, k = [], 0
    for si, (name, n) in enumerate(SECTIONS):
        groups.append((si, name, k, qs[k:k + n])); k += n
    for si, name, start, part in groups:
        para(f'{PART.get(si, "भाग स")} — {name}  (प्रश्न {start + 1}–{start + len(part)})', True, 13, before=10, after=4, keep=True)
        for j, q in enumerate(part, start + 1):
            lines = q['q'].split('\n')
            p = para(keep=True); p.paragraph_format.space_before = Pt(5)
            r = p.add_run(f'{j}. '); r.bold = True
            p.add_run(lines[0])
            for extra in lines[1:]:
                para(extra, keep=True).paragraph_format.left_indent = Pt(14)
            for o_i, o in enumerate(q['opts']):
                op = para(f'({LETTERS[o_i]}) {o}', keep=o_i < 3)
                op.paragraph_format.left_indent = Pt(14); op.paragraph_format.space_after = Pt(0)
    # उत्तरमाला — नया पृष्ठ
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    para('उत्तरमाला', True, 15, WD_ALIGN_PARAGRAPH.CENTER, 2)
    cnt = [sum(1 for q in qs if q['ans'] == x) for x in range(4)]
    para('उत्तर-वितरण: ' + ', '.join(f'{LETTERS[x]}={cnt[x]}' for x in range(4)), None, 10, WD_ALIGN_PARAGRAPH.CENTER, 6)
    for si, name, start, part in groups:
        para(f'{PART.get(si, "भाग स")} — {name}  (प्रश्न {start + 1}–{start + len(part)})', True, 12, before=8, after=3, keep=True)
        t = doc.add_table(rows=1, cols=3)
        t.style = 'Table Grid'
        for c, h in zip(t.rows[0].cells, ['क्रमांक', 'सही विकल्प', 'सही उत्तर एवं व्याख्या']):
            c.text = ''; rr = c.paragraphs[0].add_run(h); rr.bold = True
        for j, q in enumerate(part, start + 1):
            cells = t.add_row().cells
            cells[0].text = str(j); cells[1].text = f'({LETTERS[q["ans"]]})'
            cp = cells[2].paragraphs[0]
            rr = cp.add_run(q['opts'][q['ans']]); rr.bold = True
            cp.add_run('\nव्याख्या: ' + q['exp'])
        for row in t.rows:
            for w, c in zip((Emu(560000), Emu(800000), Emu(4760000)), row.cells):
                c.width = w
    doc.save(path)


def main():
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    no = int(sys.argv[1]); src = sys.argv[2]; date_txt = sys.argv[3].strip()
    iso, dpart = hindi_date(date_txt)
    qs, perr = parse_src(src)
    errs, warns = validate(qs, perr)
    for w in warns:
        print('चेतावनी:', w)
    if errs:
        for e in errs:
            print('त्रुटि:', e)
        raise SystemExit(f'\n{len(errs)} त्रुटियाँ — पहले {src} ठीक करें। कुछ भी नहीं बदला गया।')

    # 1. JSON
    secs = [list(x) for x in SECTIONS]
    js = {'secs': secs, 'qs': [{'sec': q['sec'], 'q': q['q'], 'opts': q['opts'], 'ans': q['ans'],
                                'diff': q['diff'], 'exp': q['exp'], 'fixed': False} for q in qs]}
    jpath = os.path.join(ROOT, f'cg_lecturer_set{no}.json')
    json.dump(js, open(jpath, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

    # 2. index.html
    lines, i, D = read_index()
    entry = {'no': no, 'sections': secs,
             'q': [[DIFF[q['diff']], q['q'], q['opts'], q['ans'], q['exp']] for q in qs]}
    sets = D['cg']['sets']
    n_old = len(D['q']) + sum(len(p['q']) for p in D.get('pools', {}).values()) + \
        sum(len(s['q']) for s in sets if s['no'] != no)
    pos = next((k for k, s in enumerate(sets) if s['no'] == no), None)
    if pos is None:
        sets.append(entry)
    else:
        sets[pos] = entry
    write_index(lines, i, D)
    html = open(os.path.join(ROOT, 'index.html'), encoding='utf-8').read()
    reg = sorted(s['no'] for s in sets if 26 <= s['no'] < 100)
    hi = max(reg)
    html = re.sub(r'सेट 26–\d+ व N22', f'सेट 26–{hi} व N22', html)
    html = re.sub(r'पूर्ण मॉक सेट \(09–\d+\)', f'पूर्ण मॉक सेट (09–{hi})', html)
    open(os.path.join(ROOT, 'index.html'), 'w', encoding='utf-8').write(html)

    # 3. README
    d = {k: sum(1 for q in qs if q['diff'] == k) for k in 'EMH'}
    rp = os.path.join(ROOT, 'README.md')
    rd = open(rp, encoding='utf-8').read()
    latest = (f'नवीनतम: **CG व्याख्याता (इतिहास) सेट {no} — SCERT पाठ्यक्रम आधारित** ({date_txt}; 100 प्रश्न; '
              f'{d["E"]} सरल / {d["M"]} मध्यम / {d["H"]} कठिन)।')
    rd = re.sub(r'^नवीनतम:.*$', latest, rd, flags=re.M) if re.search(r'^नवीनतम:', rd, re.M) else rd + '\n' + latest + '\n'
    rd = re.sub(r'CG सेट 09–\d+', f'CG सेट 09–{hi}', rd)
    open(rp, 'w', encoding='utf-8').write(rd)

    # 4. Word
    dname = f'CG_Lecturer_History_Set_{no}_Syllabus_Based_{dpart}.docx'
    for f in os.listdir(ROOT):  # इसी सेट की पुरानी (दूसरी तारीख वाली) फ़ाइल हटाएँ
        if f.startswith(f'CG_Lecturer_History_Set_{no}_') and f.endswith('.docx') and f != dname:
            os.remove(os.path.join(ROOT, f))
    build_docx(no, qs, date_txt, os.path.join(ROOT, dname), n_old)

    # 5. सूची और लॉग
    cnt = [sum(1 for q in qs if q['ans'] == x) for x in range(4)]
    L = [f'# CG व्याख्याता (इतिहास) — सेट {no} (SCERT पाठ्यक्रम) प्रश्न-सूची', '',
         f'बनाया गया: {date_txt} · 100 प्रश्न · {d["E"]}/{d["M"]}/{d["H"]} · '
         + ' '.join(f'{LETTERS[x]}={cnt[x]}' for x in range(4)), '',
         f'फ़ाइल: {dname} · मॉक टेस्ट: सेट {no}', '']
    k = 0
    for name, n in SECTIONS:
        L.append(f'## {name} ({n})')
        for q in qs[k:k + n]:
            k += 1
            L.append(f'{k}. {q["q"].replace(chr(10), " / ")} → ({LETTERS[q["ans"]]}) {q["opts"][q["ans"]]}')
        L.append('')
    os.makedirs(os.path.join(ROOT, 'tools', 'lists'), exist_ok=True)
    open(os.path.join(ROOT, 'tools', 'lists', f'set{no}.md'), 'w', encoding='utf-8').write('\n'.join(L))
    lp = os.path.join(ROOT, 'tools', 'log.json')
    log = json.load(open(lp, encoding='utf-8')) if os.path.exists(lp) else {}
    log[str(no)] = {'date': iso, 'date_hi': date_txt, 'docx': dname,
                    'diff': [d['E'], d['M'], d['H']], 'ans': cnt}
    json.dump(dict(sorted(log.items(), key=lambda x: int(x[0]))), open(lp, 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)

    print(f'सेट {no} जोड़ा गया ({date_txt})')
    print(f'कठिनाई: सरल {d["E"]} / मध्यम {d["M"]} / कठिन {d["H"]}')
    print('उत्तर-वितरण: ' + ' '.join(f'{LETTERS[x]}={cnt[x]}' for x in range(4)))
    print(f'Word फ़ाइल: {dname}')
    print(f'फ़ाइलें: index.html, README.md, cg_lecturer_set{no}.json, {dname}, '
          f'tools/src/set{no}.txt, tools/lists/set{no}.md, tools/log.json')


if __name__ == '__main__':
    main()
